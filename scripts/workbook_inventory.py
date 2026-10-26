from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parent.parent / 'Initial provided Info' / 'Test_files'

for p in sorted(root.glob('*.xlsx')):
    wb = load_workbook(p, read_only=True, data_only=True)
    print(f'FILE: {p.name}')
    print('  sheets:', wb.sheetnames)
    for s in wb.sheetnames:
        ws = wb[s]
        print(f'    - {s}: rows={ws.max_row}, cols={ws.max_column}, used={ws.calculate_dimension()}')
    print()
