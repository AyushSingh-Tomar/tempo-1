"""Live check (real DIAL call) of the money-relations flow on a test workbook.
Run from backend/:  venv/Scripts/python.exe .scratch/live_relations_e2e.py [file] [--no-ai]"""

import asyncio
import sys
import tempfile
import time
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, "src")
import httpx
import openpyxl

from bordereaux.config import Settings
from bordereaux.main import create_app

FIXTURES = Path("../Initial provided Info/Test_files")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
H = {"Authorization": "Bearer dev-admin1-admin"}


def read_treaty(path: Path) -> dict:
    ws = openpyxl.load_workbook(path, data_only=True)["Treaty Figures"]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, value = str(row[0] or ""), row[1]
        if value is None:
            continue
        text = str(value).strip().rstrip("%")
        try:
            num = Decimal(text) / (Decimal(100) if str(value).strip().endswith("%") or Decimal(text) > 1 else 1)
        except Exception:
            continue
        key = (
            "commission_pct" if "ommission" in name else "brokerage_pct" if "rokerage" in name else "tax_pct" if "ax" in name
            else "cession_pct" if "ession" in name or "ced" in name.lower() else "order_pct" if "rder" in name or "hare" in name else None
        )
        if key:
            out.setdefault(key, str(num))
    out.setdefault("commission_pct", "0.25")
    out.setdefault("tax_pct", "0.08")
    out.setdefault("cession_pct", "0.5")
    return out


async def main(fname: str, use_ai: bool) -> None:
    tmp = Path(tempfile.mkdtemp())
    settings = Settings(database_url=f"sqlite+aiosqlite:///{(tmp / 'live.db').as_posix()}", storage_root=tmp / "storage")
    app = create_app(settings)
    async with app.router.lifespan_context(app):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t", timeout=300) as c:
            pid = (await c.post("/api/v1/projects", json={"name": "live"}, headers=H)).json()["data"]["id"]
            f = FIXTURES / fname
            r = await c.post(f"/api/v1/projects/{pid}/processing-runs", files={"file": (f.name, f.read_bytes(), XLSX)}, headers=H)
            run_id = r.json()["data"]["run"]["id"]
            base = f"/api/v1/processing-runs/{run_id}"
            for step, kw in (("validate", {"params": {"sheet_name": "Bordereaux"}}), ("clean", {}), ("cleaning/acknowledge", {})):
                assert (await c.post(f"{base}/{step}", headers=H, **kw)).status_code == 200
            await c.put(f"{base}/mappings/target-format", json={"mode": "DEFAULT"}, headers=H)
            t = time.time()
            gen = await c.post(f"{base}/mappings/generate", params={"use_ai": str(use_ai).lower()}, headers=H)
            print(f"generate {gen.status_code} {time.time() - t:.1f}s ai={gen.json()['data']['ai']['status']}: {gen.json()['data']['ai']['message']}")
            await c.post(f"{base}/mappings/approve-exact", headers=H)
            for cd in (await c.get(f"{base}/mappings/candidates", headers=H)).json()["data"]["candidates"]:
                if cd["status"] == "PROPOSED":
                    loser = any(e.get("type") == "ambiguous_alias" or (e.get("type") == "duplicate_target_conflict" and not e.get("winner", True)) for e in cd["evidence"])
                    await c.post(f"/api/v1/mappings/{cd['id']}/{'reject' if loser else 'approve'}", json={"reason": "x"} if loser else None, headers=H)
            fin = await c.post(f"{base}/mappings/finalize", headers=H)
            print("finalize", fin.status_code, fin.text[:200] if fin.status_code != 200 else "")
            derived = (await c.get(f"{base}/treaty-figures/derived", headers=H)).json()["data"]
            body = {**read_treaty(f), "settlement_currency": derived["settlement_currency"],
                    "fx_rates": [{"currency": x["currency"], "rate": x["rate"] or "1", "source": "FILE"} for x in derived["currencies"]]}
            tr = await c.post(f"{base}/treaty-figures", json=body, headers=H)
            print("treaty", tr.status_code, tr.text[:200] if tr.status_code != 201 else "")

            ws = (await c.get(f"{base}/money-relations", headers=H)).json()["data"]
            print("numeric fields:", [(o["field"], o["label"]) for o in ws["fields"]])
            std = (await c.post(f"{base}/money-relations/standard", headers=H)).json()["data"]
            print("standard added:", std["added"], "skipped:", std["skipped"])
            for r_ in std["relations"]:
                print("   ", r_["origin"], r_["expression"])
            t = time.time()
            ai = (await c.post(f"{base}/money-relations/ai-suggest", headers=H)).json()["data"]
            print(f"ai-suggest {time.time() - t:.1f}s -> {ai['status']}: {ai['message']}")
            for s_ in ai["suggestions"]:
                print(f"    {s_['expression']}   agreement={s_['agreement']} rows={s_['rows_evaluated']} conf={s_['confidence']}  [{s_['rationale']}]")
            for s_ in ai["suggestions"][:2]:
                ar = await c.post(f"{base}/money-relations", json={**{k: s_[k] for k in ("label", "target_field", "terms", "rationale")}, "origin": "AI_SUGGESTED"}, headers=H)
                print("   add ->", ar.status_code, ar.text[:150] if ar.status_code >= 400 else "")
            run = await c.post(f"{base}/financial-checks/run", headers=H)
            print("checks", run.status_code, run.text[:200] if run.status_code != 200 else "")
            data = run.json()["data"]
            print("discrepancies:", data["discrepancies_found"])
            for rr in data["relation_results"]:
                print(f"   {rr['check_id']} {rr['origin']:12} evaluated={rr['rows_evaluated']:3} failed={rr['rows_failed']:3} skip={rr['skipped_reason']}  {rr['expression']}")
            d = (await c.get(f"{base}/discrepancies", headers=H)).json()["data"]
            print("findings by check:", {k: sum(1 for x in d["discrepancies"] if x["check_id"] == k) for k in sorted({x['check_id'] for x in d['discrepancies']})})
            print("activity:", [(i["purpose"], i["status"], i["latency_ms"]) for i in (await c.get(f"{base}/ai/activity", headers=H)).json()["data"]["interactions"]])


asyncio.run(main(next((a for a in sys.argv[1:] if not a.startswith("--")), "test_mcr_01_core_deduction.xlsx"), "--no-ai" not in sys.argv))
