"""XLSX file adapter: bytes -> sheet names / rows. Multi-sheet workbooks are
never guessed at - the caller must supply ``sheet_name`` explicitly once more
than one sheet exists (see ``ingestion.service``).

Excel stores "Mar-2024" as a real date (the 1st) and only the cell's number
format says the sender meant a month. Cells whose format shows a month and year
but no day (``mmm-yy``, ``mmmm yyyy``, ``mm/yyyy`` ...) are therefore returned
as ``"YYYY-MM"`` text, so cleaning flags them as month-precision instead of
inventing a day."""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

import openpyxl

from bordereaux.core.exceptions import IngestionError, ValidationError

_QUOTED_OR_ESCAPED = re.compile(r'"[^"]*"|\\.|\[[^\]]*\]|_.|\*.')


def is_month_only_format(number_format: str | None) -> bool:
    """True for a date format that shows a month and a year but no day or time."""
    if not number_format:
        return False
    section = number_format.split(";", 1)[0]
    code = _QUOTED_OR_ESCAPED.sub("", section).lower()
    return "y" in code and "m" in code and not any(ch in code for ch in "dhs") and "am/pm" not in code


def _cell_value(cell: Any) -> Any:
    value = cell.value
    if isinstance(value, (datetime, date)) and is_month_only_format(getattr(cell, "number_format", None)):
        return f"{value.year:04d}-{value.month:02d}"
    return value

import openpyxl

from bordereaux.core.exceptions import IngestionError, ValidationError


def list_sheet_names(content: bytes) -> list[str]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover - defensive; detection.py already checked magic bytes
        raise IngestionError(f"Could not open workbook: {exc}") from exc
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def parse_xlsx_rows(content: bytes, sheet_name: str) -> list[list[Any]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover
        raise IngestionError(f"Could not open workbook: {exc}") from exc
    try:
        if sheet_name not in wb.sheetnames:
            raise ValidationError(
                f"Sheet '{sheet_name}' not found in workbook.",
                details={"available_sheets": list(wb.sheetnames)},
            )
        ws = wb[sheet_name]
        rows = [[_cell_value(c) for c in r] for r in ws.iter_rows()]
        if not rows:
            raise IngestionError(f"Sheet '{sheet_name}' has no rows.")
        return rows
    finally:
        wb.close()
