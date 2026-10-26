"""Target format (the "final columns") a run is mapped *to*.

Default: the built-in P&C data dictionary (``bordereaux_pc_v1.yaml``, treaty
figures excluded - those come from the treaty form, not from columns), with its
built-in recognised terms.
Custom: the user supplies the column names they want, and optionally the terms
they recognise each one by. Each name is matched against the default dictionary
(exact alias, then fuzzy >= 0.90): a hit inherits the dictionary field's type /
money role / mandatory flag so downstream checks keep working; a miss becomes a
plain user-defined text field. Either way the user's own label - and every
recognised term they listed for the column - is added as an alias, so a source
header equal to any of them matches at 100%.
"""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from bordereaux.core.exceptions import ValidationError
from bordereaux.mapping.aliasing import build_field_index, exact_match, exact_matches, fuzzy_match, normalize_alias
from bordereaux.models.run import ProcessingRun
from bordereaux.schemas.loader import load_schema_definition

MAX_CUSTOM_COLUMNS = 200
MAX_TERMS_PER_COLUMN = 200
MAX_TERM_LENGTH = 200
MAX_TERMS_FILE_BYTES = 2 * 1024 * 1024
_SLUG_RE = re.compile(r"[^a-z0-9]+")


def _slug(text: str) -> str:
    return _SLUG_RE.sub("_", text.strip().lower()).strip("_") or "column"


def default_fields() -> list[dict[str, Any]]:
    definition = load_schema_definition()
    return [f for f in definition["fields"] if f.get("category") != "TREATY_FIGURES"]


def clean_terms(terms: dict[str, list[str]] | None, columns: list[str]) -> dict[str, list[str]]:
    """Validate the user's recognised terms: ``{column label: [term, ...]}``.

    Keys are matched to the final columns ignoring case/punctuation (a key that
    names no final column is an error, never silently dropped); blank and
    duplicate terms are removed; the result is keyed by the column's own label."""
    if not terms:
        return {}
    by_key = {normalize_alias(c): c for c in columns}
    cleaned: dict[str, list[str]] = {}
    unknown: list[str] = []
    for raw_key, raw_terms in terms.items():
        label = by_key.get(normalize_alias(str(raw_key)))
        if label is None:
            unknown.append(str(raw_key))
            continue
        seen = {normalize_alias(x) for x in cleaned.get(label, [])}
        out = cleaned.setdefault(label, [])
        for raw in raw_terms or []:
            if raw is None:
                continue
            term = " ".join(str(raw).split())
            key = normalize_alias(term)
            if not term or not key or key in seen:
                continue
            if len(term) > MAX_TERM_LENGTH:
                raise ValidationError(f"A recognised term for '{label}' is longer than {MAX_TERM_LENGTH} characters")
            seen.add(key)
            out.append(term)
        if len(out) > MAX_TERMS_PER_COLUMN:
            raise ValidationError(f"At most {MAX_TERMS_PER_COLUMN} recognised terms are supported per column ('{label}')")
    if unknown:
        raise ValidationError(
            "Recognised terms were given for columns that are not in the final column list",
            details={"unknown_columns": unknown},
        )
    return {label: t for label, t in cleaned.items() if t}


def build_custom_fields(columns: list[str], terms: dict[str, list[str]] | None = None) -> list[dict[str, Any]]:
    cleaned: list[str] = []
    seen: set[str] = set()
    for raw in columns:
        label = str(raw).strip()
        key = normalize_alias(label)
        if not label or not key or key in seen:
            continue
        seen.add(key)
        cleaned.append(label)
    if not cleaned:
        raise ValidationError("Provide at least one target column name")
    if len(cleaned) > MAX_CUSTOM_COLUMNS:
        raise ValidationError(f"At most {MAX_CUSTOM_COLUMNS} target columns are supported")
    user_terms = clean_terms(terms, cleaned)

    base_defs = default_fields()
    base_by_name = {f["name"]: f for f in base_defs}
    index = build_field_index({"fields": base_defs})

    fields: list[dict[str, Any]] = []
    used_names: set[str] = set()
    for label in cleaned:
        own_terms = user_terms.get(label, [])
        hit = exact_match(label, index)
        if hit is None:
            fuzzy = fuzzy_match(label, index, 0.90)
            hit = fuzzy[0] if fuzzy else None
        if hit is None and own_terms:
            # No dictionary match on the name: the user's terms may still say what it is,
            # but only when they all point at one dictionary field.
            named = {f.name: f for t in own_terms for f in exact_matches(t, index)}
            hit = next(iter(named.values())) if len(named) == 1 else None
        if hit is not None and hit.name not in used_names:
            field = dict(base_by_name[hit.name])
            field["aliases"] = [*field.get("aliases", []), label, *own_terms]
            field["display_name"] = label
            field["dictionary_match"] = hit.name
        else:
            name = _slug(label)
            while name in used_names or name in base_by_name:
                name = f"custom_{name}"
            field = {
                "name": name,
                "display_name": label,
                "type": "STRING",
                "required": "optional",
                "category": "CUSTOM",
                "aliases": [label, *own_terms],
            }
        if own_terms:
            field["user_terms"] = own_terms
        used_names.add(field["name"])
        fields.append(field)
    _let_user_terms_win(fields)
    return fields


def _let_user_terms_win(fields: list[dict[str, Any]]) -> None:
    """A term the user listed for one column is taken off the built-in synonyms of the
    other columns, so "Premium" listed under Net Premium is not also a 100% match for
    Gross Premium. If they list the same term for several columns that stays
    ambiguous and is surfaced to the reviewer as such."""
    owners: dict[str, set[str]] = {}
    for f in fields:
        for term in f.get("user_terms", []):
            owners.setdefault(normalize_alias(term), set()).add(f["name"])
    if not owners:
        return
    for f in fields:
        own_label = normalize_alias(f["display_name"])
        f["aliases"] = [
            a
            for a in f["aliases"]
            if normalize_alias(a) == own_label or f["name"] in owners.get(normalize_alias(a), {f["name"]})
        ]


def parse_columns_text(text: str) -> list[str]:
    """Split pasted text (newlines, commas, tabs or semicolons) into names."""
    return [p.strip() for p in re.split(r"[\r\n,;\t]+", text) if p.strip()]


def _read_grid(filename: str, content: bytes) -> list[list[str]]:
    """Every row of a CSV / TSV / first-sheet XLSX as stripped strings, keeping cell positions."""
    if len(content) > MAX_TERMS_FILE_BYTES:
        raise ValidationError("That file is too large to read as a list of terms")
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        import openpyxl

        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any parse failure is a user-input problem
            raise ValidationError("Could not read the Excel file") from exc
        try:
            return [["" if c is None else str(c).strip() for c in row] for row in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    if lower.endswith((".csv", ".tsv")):
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text), delimiter="\t" if lower.endswith(".tsv") else ",")
        return [[c.strip() for c in row] for row in reader]
    raise ValidationError("Unsupported file type; upload .csv, .tsv or .xlsx")


def parse_terms_file(filename: str, content: bytes) -> dict[str, Any]:
    """Read ``final columns as headers, recognised terms listed below each``.

    Returns ``{"columns": [...], "terms": {column: [term, ...]}, "skipped_terms": n}``;
    ``skipped_terms`` counts terms sitting under an empty header. Nothing is saved."""
    rows = [r for r in _read_grid(filename, content) if any(r)]
    if not rows:
        return {"columns": [], "terms": {}, "skipped_terms": 0}
    header, body = rows[0], rows[1:]
    columns: list[str] = []
    terms: dict[str, list[str]] = {}
    skipped = 0
    first_for: dict[str, str] = {}  # normalised header -> label of its first occurrence
    for j, raw in enumerate(header):
        label = raw.strip()
        values = [r[j] for r in body if j < len(r) and r[j]]
        if not label or not normalize_alias(label):
            skipped += len(values)
            continue
        label = first_for.setdefault(normalize_alias(label), label)  # a repeated header adds to the first
        if label not in columns:
            columns.append(label)
        terms.setdefault(label, []).extend(values)
    return {"columns": columns, "terms": {c: t for c, t in terms.items() if t}, "skipped_terms": skipped}


def parse_columns_file(filename: str, content: bytes) -> list[str]:
    """Header row of a CSV/XLSX file (or a one-name-per-line TXT)."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        import openpyxl

        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any parse failure is a user-input problem
            raise ValidationError("Could not read the Excel file") from exc
        ws = wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            names = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if names:
                return names
        return []
    if lower.endswith((".csv", ".txt", ".tsv")):
        text = content.decode("utf-8-sig", errors="replace")
        if lower.endswith(".txt"):
            return parse_columns_text(text)
        reader = csv.reader(io.StringIO(text), delimiter="\t" if lower.endswith(".tsv") else ",")
        for row in reader:
            names = [c.strip() for c in row if c.strip()]
            if names:
                return names
        return []
    raise ValidationError("Unsupported file type; upload .csv, .xlsx or .txt")


def effective_fields(run: ProcessingRun) -> list[dict[str, Any]]:
    """Mappable field definitions for this run (default dictionary or custom)."""
    tf = run.target_format
    if tf and tf.get("mode") == "CUSTOM" and tf.get("fields"):
        return list(tf["fields"])
    return default_fields()


def effective_definition(run: ProcessingRun) -> dict[str, Any]:
    return {"fields": effective_fields(run)}


def describe(run: ProcessingRun) -> dict[str, Any]:
    tf = run.target_format
    custom = bool(tf and tf.get("mode") == "CUSTOM" and tf.get("fields"))
    return {
        "mode": "CUSTOM" if custom else "DEFAULT",
        "selected": tf is not None,
        # The saved format this run's columns came from (a snapshot: later edits to it do not apply here).
        "template": (tf or {}).get("template") if custom else None,
        "fields": [
            {
                "name": f["name"],
                "display_name": f.get("display_name", f["name"]),
                "type": f.get("type", "STRING"),
                "required": f.get("required", "optional"),
                "category": f.get("category", ""),
                "dictionary_match": f.get("dictionary_match"),
                # Every term that maps a source header to this column, and the ones the user added.
                "recognised_terms": list(f.get("aliases", [])),
                "user_terms": list(f.get("user_terms", [])),
            }
            for f in effective_fields(run)
        ],
    }
