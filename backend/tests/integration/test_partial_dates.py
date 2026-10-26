"""Month-and-year dates (no day) and persisted period extracts, end to end through the HTTP API.

Month-only dates are kept, flagged at cleaning and in the processed file, and are
included by a quarterly / yearly (or monthly) extract that contains their month.
Each extract is stored as an immutable numbered version with audit-log entries.
"""

from __future__ import annotations

import csv
import hashlib
import io
from datetime import datetime
from pathlib import Path
from uuid import UUID

import httpx
import openpyxl
import pytest
from sqlalchemy import select

from bordereaux.models import AuditEvent, PeriodExtractArtifact, ProcessedFileArtifact
from bordereaux.models.immutability import AppendOnlyViolation
from tests.conftest import auth
from tests.integration import _pipeline as pipeline

pytestmark = pytest.mark.integration

TOKEN = pipeline.TOKEN
HEADER = ["UMR", "Policy No", "Insured Name", "Booking Date", "Gross Premium", "Commission", "IPT", "Net Premium"]
BOOKED = [
    "15/01/2024",  # 2  full date, Q1
    "Mar-2024",  # 3  month only, Q1
    "02/2024",  # 4  month only, Q1
    "May 2024",  # 5  month only, Q2
    "2024-07-09",  # 6  full date, Q3
    "Dec-2023",  # 7  month only, previous year
    None,  # 8  no date at all
]


def _save(tmp_path: Path, name: str, booked: list, formats: list[str | None] | None = None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bordereaux"
    ws.append(HEADER)
    for i, value in enumerate(booked, 1):
        ws.append([f"B{i}2024XY000{i}", f"POL-{i}", f"Insured {i}", value, 1000, 250, 120, 630])
        if formats and formats[i - 1]:
            ws.cell(row=ws.max_row, column=4).number_format = formats[i - 1]
    path = tmp_path / name
    wb.save(path)
    return path


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    return _save(tmp_path, "partial_dates.xlsx", BOOKED)


async def completed_run(client: httpx.AsyncClient, project, path: Path) -> str:
    run_id = await pipeline.map_to_treaty_step(client, project, path)
    base = f"/api/v1/processing-runs/{run_id}"
    derived = (await client.get(f"{base}/treaty-figures/derived", headers=auth(TOKEN))).json()["data"]
    body = {
        "commission_pct": "0.25", "brokerage_pct": "0.0125", "tax_pct": "0.12", "cession_pct": "1", "order_pct": "1",
        "settlement_currency": derived["settlement_currency"] or "EUR",
        "fx_rates": [{"currency": c["currency"], "rate": c["rate"] or "1", "source": "FILE"} for c in derived["currencies"]],
    }
    assert (await client.post(f"{base}/treaty-figures", json=body, headers=auth(TOKEN))).status_code == 201
    assert (await client.post(f"{base}/financial-checks/run", headers=auth(TOKEN))).status_code == 200
    assert (await client.post(f"{base}/discrepancies/proceed", headers=auth(TOKEN))).status_code == 200
    assert (await client.post(f"{base}/reports/generate", headers=auth(TOKEN))).status_code == 200
    return base


def rows_of(text: str) -> list[dict]:
    return list(csv.DictReader(io.StringIO(text)))


async def extract(client: httpx.AsyncClient, base: str, expect: int | tuple = (200, 201), **body) -> dict:
    r = await client.post(f"{base}/reports/period-extracts", json=body, headers=auth(TOKEN))
    assert r.status_code in (expect if isinstance(expect, tuple) else (expect,)), (body, r.status_code, r.text)
    return r.json()["data"] if r.status_code < 300 else r.json()


async def extracted_rows(client: httpx.AsyncClient, base: str, **body) -> list[int]:
    data = await extract(client, base, **body)
    r = await client.get(f"{base}/reports/period-extracts/{data['id']}/download", headers=auth(TOKEN))
    assert r.status_code == 200, r.text
    return [int(x["source_row_number"]) for x in rows_of(r.text)]


async def test_month_only_dates_are_flagged_kept_and_included_by_period_extracts(client: httpx.AsyncClient, project, workbook: Path):
    base = await completed_run(client, project, workbook)

    # Flagged at cleaning: one warning per month-only cell, none treated as unparseable.
    cleaning = (await client.get(f"{base}/cleaning-result", headers=auth(TOKEN))).json()["data"]["latest"]
    flagged = [w for w in cleaning["warnings"] if w["type"] == "partial_date"]
    assert [(w["row_number"], w["value"], w["normalized"]) for w in flagged] == [
        (3, "Mar-2024", "2024-03"), (4, "02/2024", "2024-02"), (5, "May 2024", "2024-05"), (7, "Dec-2023", "2023-12"),
    ]
    assert not any(w["type"] == "unparseable_date" for w in cleaning["warnings"])
    audit = (await client.get(f"{base}/download/audit_report", headers=auth(TOKEN))).json()
    assert [w["row_number"] for w in audit["cleaning"]["warnings"] if w["type"] == "partial_date"] == [3, 4, 5, 7]

    # The processed file keeps every row, month-precision values (no day) and a flag column.
    full = rows_of((await client.get(f"{base}/download/processed_file", headers=auth(TOKEN))).text)
    assert len(full) == 7
    by_row = {int(r["source_row_number"]): r for r in full}
    assert by_row[3]["booking_date"] == "2024-03" and by_row[3]["date_quality_flags"] == "booking_date=MONTH"
    assert by_row[2]["booking_date"] == "2024-01-15" and by_row[2]["date_quality_flags"] == ""
    assert by_row[8]["booking_date"] == "" and by_row[8]["date_quality_flags"] == ""

    # Yearly: all of 2024, including the three month-only dates.
    data = await extract(client, base, year=2024)
    assert (data["version_number"], data["period_label"], data["created"]) == (1, "2024", True)
    assert (data["total_rows"], data["included_rows"], data["included_partial_dates"]) == (7, 5, 3)
    assert (data["excluded_outside_period"], data["excluded_undated"]) == (1, 1)
    assert await extracted_rows(client, base, year=2024) == [2, 3, 4, 5, 6]
    # Quarterly: Mar-2024 and 02/2024 belong to Q1, May 2024 to Q2.
    assert await extracted_rows(client, base, year=2024, quarter=1) == [2, 3, 4]
    assert await extracted_rows(client, base, year=2024, quarter=2) == [5]
    assert await extracted_rows(client, base, year=2024, quarter=3) == [6]
    assert await extracted_rows(client, base, year=2023, quarter=4) == [7]
    assert await extracted_rows(client, base, year=2024, month=3) == [3]
    # The extract keeps the month-only values and their flags untouched.
    q1 = await extract(client, base, year=2024, quarter=1)
    text = (await client.get(f"{base}/reports/period-extracts/{q1['id']}/download", headers=auth(TOKEN))).text
    assert [(r["booking_date"], r["date_quality_flags"]) for r in rows_of(text)] == [
        ("2024-01-15", ""), ("2024-03", "booking_date=MONTH"), ("2024-02", "booking_date=MONTH"),
    ]


async def test_real_excel_month_formatted_cells_are_month_only_in_the_processed_file(client: httpx.AsyncClient, project, tmp_path: Path):
    """Excel stores 'Mar-24' as the date 2024-03-01; the cell's number format is the only hint."""
    booked = [datetime(2024, 2, 1), datetime(2024, 5, 1), datetime(2024, 3, 20), datetime(2023, 11, 1), datetime(2024, 3, 1)]
    formats = ["mmm-yy", "mmm-yy", "dd/mm/yyyy", "mmmm yyyy", "mm/yyyy"]
    base = await completed_run(client, project, _save(tmp_path, "excel_formats.xlsx", booked, formats))

    full = {int(r["source_row_number"]): r for r in rows_of((await client.get(f"{base}/download/processed_file", headers=auth(TOKEN))).text)}
    assert [full[n]["booking_date"] for n in range(2, 7)] == ["2024-02", "2024-05", "2024-03-20", "2023-11", "2024-03"]
    assert [full[n]["date_quality_flags"] for n in range(2, 7)] == ["booking_date=MONTH", "booking_date=MONTH", "", "booking_date=MONTH", "booking_date=MONTH"]

    # Quarterly work: the Excel month cells land in their quarter, with the dated row.
    assert await extracted_rows(client, base, year=2024, quarter=1) == [2, 4, 6]
    assert await extracted_rows(client, base, year=2024, quarter=2) == [3]
    assert await extracted_rows(client, base, year=2023, quarter=4) == [5]
    assert await extracted_rows(client, base, year=2024) == [2, 3, 4, 6]


async def test_extracts_are_versioned_immutable_and_audited(client: httpx.AsyncClient, app, db_session, project, workbook: Path):
    base = await completed_run(client, project, workbook)
    run_id = base.rsplit("/", 1)[1]

    v1 = await extract(client, base, expect=201, year=2024, quarter=1)
    again = await extract(client, base, expect=200, year=2024, quarter=1)
    assert again["id"] == v1["id"] and again["created"] is False  # unchanged result: no duplicate version
    other = await extract(client, base, expect=201, year=2024, quarter=2)
    assert other["version_number"] == 1 and other["id"] != v1["id"]  # versions are per (field, period)

    # The processed file is regenerated with different content: the next extract is version 2.
    storage = app.state.storage
    source = (await db_session.execute(select(ProcessedFileArtifact))).scalars().one()
    original_csv = await storage.download(source.storage_path)
    changed = original_csv.replace(b"POL-1,", b"POL-1X,")
    result = await storage.upload("runs/regenerated/processed_file.csv", changed, content_type="text/csv")
    db_session.add(ProcessedFileArtifact(run_id=source.run_id, storage_path=result.path, format="CSV", checksum=result.checksum,
                                         generated_at=source.generated_at.replace(year=source.generated_at.year + 1)))
    await db_session.commit()
    v2 = await extract(client, base, expect=201, year=2024, quarter=1)
    assert v2["version_number"] == 2 and v2["source_checksum"] == result.checksum and v2["source_checksum"] != v1["source_checksum"]

    # Earlier versions are never rewritten.
    d1 = await client.get(f"{base}/reports/period-extracts/{v1['id']}/download", headers=auth(TOKEN))
    d2 = await client.get(f"{base}/reports/period-extracts/{v2['id']}/download", headers=auth(TOKEN))
    assert b"POL-1X," not in d1.content and b"POL-1X," in d2.content
    assert hashlib.sha256(d1.content).hexdigest() == v1["checksum"]
    assert d1.headers["X-Extract-Version"] == "1" and d2.headers["X-Extract-Version"] == "2"
    assert "processed_file_2024Q1_booking_date_v2.csv" in d2.headers["content-disposition"]

    listing = (await client.get(f"{base}/reports/period-extracts", headers=auth(TOKEN))).json()["data"]["extracts"]
    assert sorted((e["period_label"], e["version_number"]) for e in listing) == [("2024Q1", 1), ("2024Q1", 2), ("2024Q2", 1)]

    # Audit trail: one entry per generated version, one per download.
    events = (await db_session.execute(select(AuditEvent).where(AuditEvent.event_type.like("reports.period_extract_%")))).scalars().all()
    generated = [e for e in events if e.event_type == "reports.period_extract_generated"]
    assert len(generated) == 3 and {e.after_state["period"] for e in generated} == {"2024Q1", "2024Q2"}
    assert {e.target_id.hex for e in generated} == {v1["id"].replace("-", ""), other["id"].replace("-", ""), v2["id"].replace("-", "")}
    assert all(e.actor_id is not None and e.after_state["checksum"] for e in generated)
    assert len([e for e in events if e.event_type == "reports.period_extract_downloaded"]) == 2

    # The table itself refuses edits.
    row = (await db_session.execute(select(PeriodExtractArtifact).where(PeriodExtractArtifact.version_number == 1))).scalars().first()
    row.included_rows = 999
    with pytest.raises(AppendOnlyViolation):
        await db_session.flush()
    await db_session.rollback()

    # A stored file that no longer matches its recorded checksum is refused, not served.
    stored_path = (await db_session.get(PeriodExtractArtifact, UUID(v1["id"]))).storage_path
    await storage.upload(stored_path, b"tampered", overwrite=True)
    bad = await client.get(f"{base}/reports/period-extracts/{v1['id']}/download", headers=auth(TOKEN))
    assert bad.status_code == 500 and "checksum" in bad.text


async def test_period_extract_validation_and_errors(client: httpx.AsyncClient, project, workbook: Path):
    base = await completed_run(client, project, workbook)
    for body in (
        {"year": 2024, "quarter": 5}, {"year": 2024, "quarter": 1, "month": 1}, {"year": 1800}, {"quarter": 1},
        {"year": 2024, "period_field": "nope"}, {"year": 2024, "period_field": "gross_premium"},
    ):
        r = await client.post(f"{base}/reports/period-extracts", json=body, headers=auth(TOKEN))
        assert r.status_code in (400, 422), (body, r.status_code)

    # The old query-string filter fails loudly instead of silently returning the whole file.
    r = await client.get(f"{base}/download/processed_file", params={"year": 2024}, headers=auth(TOKEN))
    assert r.status_code == 400 and "period-extracts" in r.text

    unknown = "00000000-0000-0000-0000-000000000000"
    assert (await client.get(f"{base}/reports/period-extracts/{unknown}/download", headers=auth(TOKEN))).status_code == 404

    # An extract of another run is not reachable through this run's URL.
    own = await extract(client, base, year=2024)
    other_base = base.rsplit("/", 1)[0] + f"/{unknown}"
    assert (await client.get(f"{other_base}/reports/period-extracts/{own['id']}/download", headers=auth(TOKEN))).status_code == 404


async def test_extract_needs_generated_reports(client: httpx.AsyncClient, project, workbook: Path):
    run_id = await pipeline.map_to_treaty_step(client, project, workbook)
    r = await client.post(f"/api/v1/processing-runs/{run_id}/reports/period-extracts", json={"year": 2024}, headers=auth(TOKEN))
    assert r.status_code == 409
