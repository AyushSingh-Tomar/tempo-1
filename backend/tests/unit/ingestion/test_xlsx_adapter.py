from __future__ import annotations

import io
from datetime import datetime

import openpyxl
import pytest

from bordereaux.ingestion.xlsx_adapter import is_month_only_format, parse_xlsx_rows

pytestmark = pytest.mark.unit


@pytest.mark.parametrize(
    "fmt",
    ["mmm-yy", "mmm-yyyy", "mmmm yyyy", "mm/yyyy", "m/yyyy", "yyyy-mm", "mmm yy", "[$-409]mmm-yy;@", r"mmm\-yy", '"Period" mmm yyyy'],
)
def test_month_and_year_formats_are_month_only(fmt):
    assert is_month_only_format(fmt) is True


@pytest.mark.parametrize(
    "fmt",
    [None, "", "General", "0.00", "dd/mm/yyyy", "d-mmm-yy", "yyyy-mm-dd", "mm:ss", "h:mm", "mmm-yy hh:mm", "yyyy", "mmm", "dd-mmm-yyyy;@"],
)
def test_everything_else_keeps_its_day(fmt):
    assert is_month_only_format(fmt) is False


def _workbook(cells: list[tuple[str, datetime, str]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Booked"])
    for _, value, fmt in cells:
        ws.append([value])
        ws.cell(row=ws.max_row, column=1).number_format = fmt
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_month_formatted_date_cells_come_back_as_year_month_text_and_other_dates_are_untouched():
    content = _workbook(
        [
            ("a", datetime(2024, 3, 1), "mmm-yy"),
            ("b", datetime(2024, 5, 1), "mmmm yyyy"),
            ("c", datetime(2024, 7, 9), "dd/mm/yyyy"),
            ("d", datetime(2023, 12, 1), "mm/yyyy"),
        ]
    )
    rows = parse_xlsx_rows(content, "S")
    assert [r[0] for r in rows] == ["Booked", "2024-03", "2024-05", datetime(2024, 7, 9), "2023-12"]
